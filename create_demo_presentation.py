from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "compliance_iq_bot_demo.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(0x0D, 0x1B, 0x3A)
TEAL = RGBColor(0x00, 0x8A, 0x9E)
LIGHT = RGBColor(0xF4, 0xF7, 0xFA)
GRAY = RGBColor(0x5B, 0x6B, 0x7A)
DARK = RGBColor(0x1F, 0x29, 0x33)
ACCENT = RGBColor(0x2E, 0x86, 0xC1)


def add_background(slide, color):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_title(slide, title, subtitle=None):
    text_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(12.0), Inches(1.1))
    tf = text_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(14)
        p2.font.color.rgb = RGBColor(0xE8, 0xF0, 0xF6)
        p2.level = 1


def add_bullets(slide, title, bullets, left=0.7, top=1.3, width=5.7, height=4.8, title_color=RGBColor(0xFF, 0xFF, 0xFF)):
    title_box = slide.shapes.add_textbox(Inches(left), Inches(top - 0.1), Inches(width), Inches(0.4))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = title_color

    body_box = slide.shapes.add_textbox(Inches(left), Inches(top + 0.35), Inches(width), Inches(height))
    tf = body_box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(18)
        p.font.color.rgb = DARK
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)


def add_box(slide, left, top, width, height, title, body, fill=RGBColor(0xFF, 0xFF, 0xFF), title_color=ACCENT):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(0xD6, 0xDE, 0xE3)
    shape.line.width = Pt(1.2)

    tb = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.15), Inches(width - 0.4), Inches(height - 0.4))
    tf = tb.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = title_color
    p.space_after = Pt(4)
    p2 = tf.add_paragraph()
    p2.text = body
    p2.font.size = Pt(12)
    p2.font.color.rgb = GRAY
    p2.level = 0

# Slide 1: Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, NAVY)
add_title(slide, "Compliance IQ Bot", "AI-powered regulatory compliance assistant for fast, evidence-based answers")

# Title area card
card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(2.1), Inches(11.5), Inches(2.5))
card.fill.solid()
card.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
card.line.color.rgb = RGBColor(0xD6, 0xDE, 0xE3)
card.line.width = Pt(1.0)

text_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.5), Inches(10.8), Inches(1.2))
text_frame = text_box.text_frame
text_frame.clear()
p = text_frame.paragraphs[0]
p.text = "Capstone project demo: upload regulation documents, ask compliance questions, and receive grounded answers with citations."
p.font.size = Pt(20)
p.font.color.rgb = DARK
p.alignment = PP_ALIGN.CENTER

# Add small footer
author_box = slide.shapes.add_textbox(Inches(0.9), Inches(6.5), Inches(6.6), Inches(0.4))
text_frame = author_box.text_frame
text_frame.clear()
p = text_frame.paragraphs[0]
p.text = "Built with FastAPI, Streamlit, LangChain, and OpenAI-powered agents"
p.font.size = Pt(14)
p.font.color.rgb = RGBColor(0xE8, 0xF0, 0xF6)

# Slide 2: Problem
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, LIGHT)
add_title(slide, "Why this project matters", "Regulatory review is often slow, manual, and fragmented")
add_bullets(
    slide,
    "Challenges in compliance work",
    [
        "Teams spend hours searching across long regulatory PDFs and policy documents.",
        "Important obligations can be missed when information is scattered across sources.",
        "Compliance teams need quick, traceable answers instead of generic summaries.",
    ],
    left=0.9,
    top=1.6,
    width=5.8,
    height=4.3,
)

add_box(
    slide,
    7.0,
    1.8,
    5.2,
    2.2,
    "Business impact",
    "Reduces manual effort, speeds up response time, and improves audit readiness with evidence-backed answers.",
)
add_box(
    slide,
    7.0,
    4.3,
    5.2,
    2.1,
    "Target users",
    "Compliance officers, legal teams, risk managers, and operations leaders who need instant policy guidance.",
)

# Slide 3: Solution overview
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, NAVY)
add_title(slide, "What the bot does", "A conversational assistant for document-grounded compliance intelligence")
add_bullets(
    slide,
    "Core capabilities",
    [
        "Upload a compliance PDF through a simple interface.",
        "Ask natural-language questions like “What are the policy requirements?”",
        "Receive concise answers with citations to the uploaded document.",
        "Use retrieval and agent-based reasoning to ground responses in source material.",
    ],
    left=0.9,
    top=1.6,
    width=6.0,
    height=4.0,
)
add_box(slide, 7.3, 1.8, 5.0, 1.8, "Input", "PDF upload and a user question")
add_box(slide, 7.3, 3.9, 5.0, 1.8, "Output", "Structured answer + supporting citations")

# Slide 4: Architecture
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, LIGHT)
add_title(slide, "System architecture", "Frontend, API, services, and AI reasoning layers")
add_box(slide, 0.7, 1.8, 2.7, 1.6, "1. UI", "Streamlit chat experience")
add_box(slide, 3.7, 1.8, 2.7, 1.6, "2. API", "FastAPI endpoints for upload and query")
add_box(slide, 6.7, 1.8, 2.7, 1.6, "3. Services", "Document handling and query processing")
add_box(slide, 9.7, 1.8, 2.7, 1.6, "4. Agent", "LangChain agent with tools and prompts")
add_box(slide, 2.4, 4.2, 8.5, 1.8, "Flow", "Upload PDF → process document → retrieve evidence → generate grounded response with citations")

# Slide 5: Demo workflow
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, NAVY)
add_title(slide, "How the demo works", "A simple end-to-end flow for users")
add_bullets(
    slide,
    "Demo steps",
    [
        "Step 1: User uploads a PDF through the Streamlit UI.",
        "Step 2: The backend saves the file and prepares it for retrieval.",
        "Step 3: The user asks a compliance question in natural language.",
        "Step 4: The agent searches the document and returns a grounded answer with citations.",
    ],
    left=0.9,
    top=1.6,
    width=7.0,
    height=4.0,
)
add_box(slide, 8.4, 1.8, 4.0, 2.0, "Example prompt", "What are the compliance requirements mentioned in this document?")
add_box(slide, 8.4, 4.2, 4.0, 1.7, "Expected result", "A concise answer with source-backed evidence")

# Slide 6: Project stack
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, LIGHT)
add_title(slide, "Technology stack", "Modern Python stack for AI + web app delivery")
add_bullets(
    slide,
    "Key technologies",
    [
        "Python for the core application and orchestration.",
        "FastAPI for robust backend APIs.",
        "Streamlit for a lightweight chat-based UI.",
        "LangChain and OpenAI models for retrieval and reasoning.",
        "Pydantic models for structured request and response payloads.",
    ],
    left=0.9,
    top=1.6,
    width=6.0,
    height=4.3,
)
add_box(slide, 7.2, 1.8, 5.0, 1.7, "Backend", "FastAPI, services, routes, and agent integration")
add_box(slide, 7.2, 3.9, 5.0, 1.8, "AI layer", "Retrieval tools, prompts, and structured LLM responses")

# Slide 7: Benefits and next steps
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, NAVY)
add_title(slide, "Benefits and next steps", "A strong foundation for enterprise compliance copilots")
add_bullets(
    slide,
    "Value delivered",
    [
        "Accelerates compliance investigations and policy lookup.",
        "Provides transparent, citation-based answers instead of unsupported guesses.",
        "Creates a reusable foundation for future compliance automation.",
    ],
    left=0.9,
    top=1.6,
    width=6.0,
    height=3.8,
)
add_bullets(
    slide,
    "Future enhancements",
    [
        "Support more document formats such as DOCX and TXT.",
        "Add authentication, role-based access, and multi-user deployment.",
        "Expand to multiple regulatory sources and enterprise knowledge bases.",
    ],
    left=7.3,
    top=1.6,
    width=5.2,
    height=3.8,
)

# Slide 8: Thank you
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, LIGHT)
add_title(slide, "Thank you", "Questions? Ready for a live demo")

# Small contact box
box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(2.3), Inches(2.6), Inches(8.7), Inches(2.3))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
box.line.color.rgb = RGBColor(0xD6, 0xDE, 0xE3)
box.line.width = Pt(1.0)
text_box = slide.shapes.add_textbox(Inches(2.7), Inches(3.1), Inches(7.9), Inches(1.3))
text_frame = text_box.text_frame
text_frame.clear()
p = text_frame.paragraphs[0]
p.text = "Demo file created successfully: compliance_iq_bot_demo.pptx"
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = DARK
p.alignment = PP_ALIGN.CENTER

prs.save(OUTPUT)
print(f"Presentation created: {OUTPUT}")
